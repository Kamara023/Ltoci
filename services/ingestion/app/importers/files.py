"""Importeurs de fichiers CSV / Excel / JSON — parsing PUR (sans DB).

Format de ligne attendu (mapping de colonnes configurable via
`data_sources.config.columns`, défauts ci-dessous) :

  date     : 'YYYY-MM-DD' ou 'JJ/MM/AAAA'
  tirage   : nom ou code du tirage ('Réveil', 'reveil'…)
  gagnants : 5 numéros séparés par '-', ',', ';' ou espaces
  machine  : idem, optionnel

Les lignes invalides sont collectées avec leur raison — jamais insérées.
"""

from __future__ import annotations

import csv
import io
import json as jsonlib
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from app.parsers.lonaci_api import ParsedDraw, slugify

DEFAULT_COLUMNS = {"date": "date", "draw": "tirage", "winning": "gagnants", "machine": "machine"}


@dataclass
class ImportResult:
    draws: list[ParsedDraw] = field(default_factory=list)
    rejected: list[dict] = field(default_factory=list)  # {line, reason}

    @property
    def found(self) -> int:
        return len(self.draws) + len(self.rejected)


def parse_flexible_numbers(value: object) -> list[int] | None:
    """'1-2-3-4-5', '1, 2, 3, 4, 5', '1 2 3 4 5' -> [1,2,3,4,5]."""
    if value is None:
        return None
    tokens = [t for t in re.split(r"[-,;\s]+", str(value).strip()) if t]
    if not tokens or not all(t.isdigit() for t in tokens):
        return None
    return [int(t) for t in tokens]


def parse_flexible_date(value: object) -> date | None:
    text = str(value).strip() if value is not None else ""
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    if isinstance(value, datetime):
        return value.date()
    return None


def _row_to_draw(row: dict, columns: dict, line_no: int, result: ImportResult) -> None:
    normalized = {str(k).strip().lower(): v for k, v in row.items() if k is not None}
    draw_date = parse_flexible_date(normalized.get(columns["date"]))
    draw_name = str(normalized.get(columns["draw"]) or "").strip()
    winning = parse_flexible_numbers(normalized.get(columns["winning"]))
    machine = parse_flexible_numbers(normalized.get(columns.get("machine", "machine")))

    if draw_date is None:
        result.rejected.append({"line": line_no, "reason": "date illisible"})
        return
    if not draw_name:
        result.rejected.append({"line": line_no, "reason": "tirage manquant"})
        return
    if winning is None:
        result.rejected.append({"line": line_no, "reason": "numéros gagnants illisibles"})
        return

    result.draws.append(
        ParsedDraw(
            draw_name=draw_name,
            draw_code=slugify(draw_name),
            category="standard",
            draw_date=draw_date,
            winning=winning,
            machine=machine,
            raw={"line": line_no, **{k: str(v) for k, v in normalized.items()}},
        )
    )


def parse_csv(content: bytes, columns: dict | None = None) -> ImportResult:
    cols = {**DEFAULT_COLUMNS, **(columns or {})}
    result = ImportResult()
    text = content.decode("utf-8-sig")
    sample = text[:2048]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    for i, row in enumerate(reader, start=2):  # ligne 1 = en-têtes
        _row_to_draw(row, cols, i, result)
    return result


def parse_excel(content: bytes, columns: dict | None = None) -> ImportResult:
    from openpyxl import load_workbook

    cols = {**DEFAULT_COLUMNS, **(columns or {})}
    result = ImportResult()
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        headers = [str(h).strip().lower() if h is not None else "" for h in next(rows)]
    except StopIteration:
        return result
    for i, values in enumerate(rows, start=2):
        row = dict(zip(headers, values, strict=False))
        if all(v is None for v in values):
            continue
        _row_to_draw(row, cols, i, result)
    return result


def parse_json(content: bytes, columns: dict | None = None) -> ImportResult:
    cols = {**DEFAULT_COLUMNS, **(columns or {})}
    result = ImportResult()
    try:
        data = jsonlib.loads(content.decode("utf-8"))
    except (ValueError, UnicodeDecodeError):
        result.rejected.append({"line": 0, "reason": "JSON invalide"})
        return result
    if not isinstance(data, list):
        result.rejected.append({"line": 0, "reason": "le JSON doit être une liste d'objets"})
        return result
    for i, row in enumerate(data, start=1):
        if not isinstance(row, dict):
            result.rejected.append({"line": i, "reason": "élément non-objet"})
            continue
        _row_to_draw(row, cols, i, result)
    return result


PARSERS_BY_KIND = {"CSV": parse_csv, "EXCEL": parse_excel, "JSON": parse_json}
